from flask import Flask, render_template, request, redirect, url_for, flash, jsonify
from openpyxl import load_workbook
from datetime import datetime

app = Flask(__name__)
app.secret_key = 'your_secret_key'

# Konfigurasi file Excel
EXCEL_FILE = r'E:\App-Development\visitor-log\data\visitor-log.xlsx' #Direktori file excel yang digunakan untuk menyimpan data visitor
PENGUNJUNG_SHEET = 'Pengunjung' #Nama sheet di file excel yang digunakan untuk menyimpan data visitor
KUNJUNGAN_SHEET = 'Kunjungan' #Nama sheet di file excel yang digunakan untuk merekam data kunjungan visitor

@app.route("/")
def index():
    return redirect(url_for("home"))

@app.route("/visilog", methods=["POST", "GET"])
def home():
    if request.method == 'POST':
        id_pengunjung = request.form.get('id_pengunjung')

        try:
            # Load workbook dan sheet
            wb = load_workbook(EXCEL_FILE)
            pengunjung_sheet = wb[PENGUNJUNG_SHEET]
            kunjungan_sheet = wb[KUNJUNGAN_SHEET]

            # Cari data pengunjung berdasarkan ID
            pengunjung = None
            for row in pengunjung_sheet.iter_rows(min_row=2, values_only=True):
                if str(row[0]) == id_pengunjung:  # Kolom pertama adalah ID
                    pengunjung = row
                    break

            if pengunjung:
                # Rekam kunjungan
                TANGGAL_KUNJUNGAN = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                kunjungan_sheet.append([pengunjung[0], pengunjung[1], pengunjung[2], TANGGAL_KUNJUNGAN])
                wb.save(EXCEL_FILE)

                return jsonify({
                    'success': True,
                    'nama': pengunjung[1],
                    'redirect': url_for('home')
                })

            else:
                return jsonify({
                    'success': False,
                    'redirect': url_for('home')
                })
            
        except Exception as e:
            flash(f"Terjadi kesalahan: {e}", 'danger')

        return redirect(url_for('home'))

    return render_template('visitor-log.html')

if __name__ == '__main__':
    app.run(debug = True, host = '0.0.0.0', port = 5555)
